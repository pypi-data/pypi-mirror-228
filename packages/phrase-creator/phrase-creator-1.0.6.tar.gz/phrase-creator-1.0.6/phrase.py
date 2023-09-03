#!/usr/bin/env python3
#!/usr/bin/env python
import pandas as pd
from openpyxl import load_workbook
import json
import datetime
import argparse


# Get the current date
current_date = datetime.date.today()
# Format the current date as DD/MM
formatted_date = datetime.datetime.now().strftime("%d/%m")#datetime.datetime.now().replace(hour=0,minute=0,second=0,microsecond=0)
base_path="./"

def read_and_create_change_set(file_name=f"{base_path}Phrase.xlsx", sheet_name="COB-R5", filter_column=formatted_date):
    workbook = load_workbook(file_name.strip(),data_only=True)
    worksheet = workbook[sheet_name]
    data =[]
    for row in worksheet: 
        if worksheet.row_dimensions[row[0].row].hidden == False:
            row_values = [cell.value for cell in row]
            data.append(row_values)

    df = pd.DataFrame(data[1:], columns=data[0])
    df = df[(df[filter_column] == '/')]
    #Read Column Named Statement and concat 
    concat_text="["
    for i in range(0, df['Statement'].size):
        if str(df['Statement'].iloc[i]).strip() == '' or str(df['Statement'].iloc[i]).strip() == 'None':
            continue
        concat_text+=str(df['Statement'].iloc[i]).replace('_x000D_', '')
        if i < df['Statement'].size-1:
            concat_text+=","
    concat_text+= "]"
    return concat_text

    #Write into file

def create_template(result, authorName, today):
    json_template = '''{{
    "databaseChangeLog": [
        {{
            "changeSet": {{
                "id": "commonapply-XX",
                "author": "{authorName}",
                "comment": "update phase as R5 commented {series}",
                "tagDatabase": {{
                    "tag": "update-phase-config-commonapply-XX"
                }},
                "changes": [
                    {{
                        "runCommand": {{
                            "command": {{
                                "$rawJson": {{
                                    "update": "phrases_config_migration",
                                    "updates": {result},
                                    "ordered": true
                                }}
                            }}
                        }}
                    }}
                ]
            }}
        }}
    ]
}}'''

    # Replace placeholders with actual values
    formatted_json = json_template.format(authorName=authorName, series=today,result=result )
    return formatted_json

def write_file(result):
    with open("./db.changeset.json", 'w') as file:
        file.write(result)
def main():
    parser = argparse.ArgumentParser(description="Pharse Convertor from Volk")
    parser.add_argument('--file-name',help='Path to pharse file', default='./Phrase.xlsx')
    parser.add_argument('--sheet-name', help='sheet name', default='COB-R5')
    parser.add_argument('--filter-column', help='filter column', default=formatted_date)
    args = parser.parse_args()
    print(args)

    concat_text= read_and_create_change_set(args.file_name, args.sheet_name, args.filter_column)
    authorName = input("What is your name: ")
    if authorName is None or len(authorName) == 0:
        authorName = "commonapply"
    result = create_template(concat_text, authorName,formatted_date)
    write_file(result)

if __name__ == "__main__":
    main()